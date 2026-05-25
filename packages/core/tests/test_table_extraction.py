"""End-to-end test for native PPTX table extraction.

Builds a tiny synthetic .pptx with one table, runs the parser, and verifies
the resulting shape dict has the expected cell matrix and header flag.
"""
from __future__ import annotations

import pytest

pytest.importorskip("pptx")


def _build_deck_with_table(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title-only layout
    rows, cols = 3, 2
    left = top = Inches(1)
    width = Inches(4)
    height = Inches(2)
    tbl_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    tbl = tbl_shape.table
    tbl.cell(0, 0).text = "Region"
    tbl.cell(0, 1).text = "Revenue"
    tbl.cell(1, 0).text = "West"
    tbl.cell(1, 1).text = "100"
    tbl.cell(2, 0).text = "East"
    tbl.cell(2, 1).text = "200"
    prs.save(str(path))


def test_table_extraction_returns_cell_matrix(tmp_path):
    pptx_path = tmp_path / "deck.pptx"
    _build_deck_with_table(pptx_path)
    from packages.core.ppt_parser import parse_pptx
    slides = parse_pptx(str(pptx_path))
    assert len(slides) == 1
    tables = [s for s in slides[0]["shapes"] if s.get("type") == "TABLE"]
    assert len(tables) == 1
    t = tables[0]
    assert t["cells"] == [
        ["Region", "Revenue"],
        ["West", "100"],
        ["East", "200"],
    ]
    assert t["first_row_is_header"] is True
    # Tables don't carry text_runs (so the SHAPE_LABEL loop won't pick them up)
    assert t["text_runs"] == []


def test_table_extraction_is_robust_to_empty_table(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tbl_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(3), Inches(2))
    # leave all cells empty
    pptx_path = tmp_path / "empty.pptx"
    prs.save(str(pptx_path))

    from packages.core.ppt_parser import parse_pptx
    slides = parse_pptx(str(pptx_path))
    tables = [s for s in slides[0]["shapes"] if s.get("type") == "TABLE"]
    assert len(tables) == 1
    assert tables[0]["cells"] == [["", ""], ["", ""]]
    assert tables[0]["first_row_is_header"] is False  # no non-empty header cells
