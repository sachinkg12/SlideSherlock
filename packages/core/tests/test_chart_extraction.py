"""End-to-end test for native PPTX chart extraction.

Builds a small synthetic .pptx with one modern Insert→Chart object and verifies
the parser produces a CHART shape dict with title, categories, series, and values.

OLE-blob chart extraction is exercised separately via the real Drendel3 fixture
in the regression run; openpyxl/xlrd construction of a synthetic OLE-embedded
Excel inside a PPTX is hairy and not worth a unit-test investment.
"""
from __future__ import annotations

import pytest

pytest.importorskip("pptx")


def _build_deck_with_chart(path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    data = CategoryChartData()
    data.categories = ["Q1", "Q2", "Q3"]
    data.add_series("Revenue", (100, 150, 175))
    data.add_series("Cost", (60, 80, 95))
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(1), Inches(1), Inches(6), Inches(4), data,
    )
    chart_shape.chart.has_title = True
    chart_shape.chart.chart_title.text_frame.text = "Quarterly Performance"
    prs.save(str(path))


def test_chart_extraction_native_modern(tmp_path):
    pptx_path = tmp_path / "deck.pptx"
    _build_deck_with_chart(pptx_path)
    from packages.core.ppt_parser import parse_pptx
    slides = parse_pptx(str(pptx_path))
    charts = [sh for sh in slides[0]["shapes"] if sh.get("type") == "CHART"]
    assert len(charts) == 1
    c = charts[0]
    assert c["chart_title"] == "Quarterly Performance"
    assert c["categories"] == ["Q1", "Q2", "Q3"]
    assert c["source_kind"] == "native"
    labels = {s["label"] for s in c["series"]}
    assert labels == {"Revenue", "Cost"}
    revenue = next(s for s in c["series"] if s["label"] == "Revenue")
    assert revenue["values"] == [100, 150, 175]
    # text_runs is empty so the SHAPE_LABEL loop won't pick this up
    assert c["text_runs"] == []
